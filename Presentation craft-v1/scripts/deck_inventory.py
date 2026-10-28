#!/usr/bin/env python3
"""Structural inventory of a .pptx deck for presentation-quality-review.

Usage:
    python deck_inventory.py deck.pptx            # human-readable summary
    python deck_inventory.py deck.pptx --json     # full JSON

Reports facts, not judgments: per-slide titles, body word counts, font
families and sizes, explicit colors, images (with effective resolution),
charts (type, 3D), tables, animations, layouts, notes and alt text.

Font sizes/families inherited from the master are not resolvable without a
full style cascade; explicit values are reported and inherited runs are
counted separately. Always confirm visually on rendered slide images.
"""
import json
import sys
from collections import Counter

try:
    from pptx import Presentation
    from pptx.util import Emu
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    from pptx.enum.dml import MSO_FILL
except ImportError:
    sys.exit("python-pptx missing: pip install python-pptx --break-system-packages")

NS_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


def theme_fonts(prs):
    """Return {'+mj-lt': major latin, '+mn-lt': minor latin} from the first theme."""
    out = {}
    try:
        master = prs.slide_masters[0]
        for rel in master.part.rels.values():
            if rel.reltype.endswith("/theme"):
                from lxml import etree
                root = etree.fromstring(rel.target_part.blob)
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                maj = root.find(".//a:majorFont/a:latin", ns)
                mn = root.find(".//a:minorFont/a:latin", ns)
                if maj is not None:
                    out["+mj-lt"] = maj.get("typeface")
                if mn is not None:
                    out["+mn-lt"] = mn.get("typeface")
    except Exception:
        pass
    return out


def is_title(shape):
    try:
        return shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES
    except Exception:
        return False


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def has_alt(shape):
    """True if the shape has meaningful alt text (not empty, not a bare file name)."""
    vals = shape._element.xpath(".//p:cNvPr/@descr")
    d = (vals[0] if vals else "").strip()
    return bool(d) and not d.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".svg", ".emf", ".wmf"))


def color_hex(font_or_fill):
    try:
        c = font_or_fill.color
        if c and c.type is not None and c.rgb is not None:
            return "#" + str(c.rgb)
    except Exception:
        pass
    return None


def text_stats(tf, tfonts, rec, is_title_shape):
    words = 0
    for p in tf.paragraphs:
        for r in p.runs:
            t = r.text.strip()
            if not t:
                continue
            words += len(t.split())
            f = r.font
            name = f.name or (p.font.name if p.font else None)
            if name in tfonts:
                name = tfonts[name]
            if name:
                rec["fonts"][name] += 1
            else:
                rec["inherited_font_runs"] += 1
            size = f.size or (p.font.size if p.font else None)
            if size:
                pt = round(Emu(size).pt, 1)
                rec["sizes"].append(pt)
                if is_title_shape:
                    rec["title_sizes"].append(pt)
            else:
                rec["inherited_size_runs"] += 1
            col = color_hex(f)
            if col:
                rec["text_colors"][col] += 1
    return words


def analyze(path):
    prs = Presentation(path)
    tfonts = theme_fonts(prs)
    sw, shh = prs.slide_width, prs.slide_height
    slides = []
    deck_fonts, deck_colors, deck_fills = Counter(), Counter(), Counter()
    all_title_sizes = []

    for i, slide in enumerate(prs.slides, 1):
        rec = {
            "slide": i,
            "layout": slide.slide_layout.name,
            "title": None,
            "body_words": 0,
            "fonts": Counter(),
            "sizes": [],
            "title_sizes": [],
            "inherited_font_runs": 0,
            "inherited_size_runs": 0,
            "text_colors": Counter(),
            "fill_colors": Counter(),
            "images": [],
            "charts": [],
            "tables": [],
            "shapes": 0,
            "off_slide_shapes": 0,
            "has_animation": slide._element.find(f"{NS_P}timing") is not None,
            "has_transition": slide._element.find(f"{NS_P}transition") is not None,
            "notes_words": 0,
        }
        for sh in iter_shapes(slide.shapes):
            rec["shapes"] += 1
            try:
                if sh.left is not None and (sh.left < 0 or sh.top < 0 or
                                            sh.left + sh.width > sw or sh.top + sh.height > shh):
                    rec["off_slide_shapes"] += 1
            except Exception:
                pass
            title_shape = is_title(sh)
            if sh.has_text_frame:
                w = text_stats(sh.text_frame, tfonts, rec, title_shape)
                if title_shape:
                    rec["title"] = sh.text_frame.text.strip().replace("\n", " ")
                else:
                    rec["body_words"] += w
            try:
                if sh.fill.type == MSO_FILL.SOLID:
                    rec["fill_colors"]["#" + str(sh.fill.fore_color.rgb)] += 1
            except Exception:
                pass  # theme colors / no fill: not resolvable without cascade
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = {"alt_text": has_alt(sh)}
                try:
                    px_w, px_h = sh.image.size
                    disp_in = Emu(sh.width).inches
                    img["px"] = f"{px_w}x{px_h}"
                    img["effective_ppi"] = round(px_w / disp_in) if disp_in else None
                    img["aspect_distortion"] = round(
                        abs((px_w / px_h) / (sh.width / sh.height) - 1), 3)
                except Exception:
                    pass
                rec["images"].append(img)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                cxml = sh.chart.part.blob.decode("utf-8", "ignore")
                try:
                    ct = str(sh.chart.chart_type)
                except Exception:
                    ct = "unknown"
                rec["charts"].append({
                    "type": ct,
                    "is_3d": "3DChart" in cxml,
                    "has_legend": sh.chart.has_legend,
                    "series": sum(len(list(p.series)) for p in sh.chart.plots),
                    "alt_text": has_alt(sh),
                    "categories": len(sh.chart.plots[0].categories) if len(sh.chart.plots) else None,
                })
            if getattr(sh, "has_table", False) and sh.has_table:
                t = sh.table
                rec["tables"].append({"rows": len(t.rows), "cols": len(t.columns)})
                for row in t.rows:
                    for cell in row.cells:
                        rec["body_words"] += text_stats(cell.text_frame, tfonts, rec, False)
        if slide.has_notes_slide:
            rec["notes_words"] = len(slide.notes_slide.notes_text_frame.text.split())

        deck_fonts.update(rec["fonts"])
        deck_colors.update(rec["text_colors"])
        deck_fills.update(rec["fill_colors"])
        all_title_sizes += rec["title_sizes"]
        rec["min_size"] = min(rec["sizes"]) if rec["sizes"] else None
        for k in ("fonts", "text_colors", "fill_colors"):
            rec[k] = dict(rec[k])
        rec["sizes"] = sorted(set(rec["sizes"]))
        rec["title_sizes"] = sorted(set(rec["title_sizes"]))
        slides.append(rec)

    all_sizes = sorted({s for r in slides for s in r["sizes"]})
    return {
        "file": path,
        "slide_count": len(slides),
        "slide_size_in": f"{Emu(sw).inches:.2f}x{Emu(shh).inches:.2f}",
        "theme_fonts": tfonts,
        "deck": {
            "font_families": dict(deck_fonts.most_common()),
            "font_sizes_pt": all_sizes,
            "title_sizes_pt": sorted(set(all_title_sizes)),
            "text_colors": dict(deck_colors.most_common()),
            "fill_colors": dict(deck_fills.most_common()),
            "layouts": dict(Counter(r["layout"] for r in slides).most_common()),
            "slides_with_animation": [r["slide"] for r in slides if r["has_animation"]],
            "slides_with_transition": [r["slide"] for r in slides if r["has_transition"]],
            "slides_without_title": [r["slide"] for r in slides if not r["title"]],
            "body_words_per_slide": {r["slide"]: r["body_words"] for r in slides},
        },
        "slides": slides,
    }


def summary(d):
    D = d["deck"]
    out = [f"# Inventory: {d['file']}",
           f"{d['slide_count']} slides · {d['slide_size_in']} in · theme fonts {d['theme_fonts'] or 'n/a'}",
           "",
           f"Font families (explicit runs): {D['font_families'] or 'all inherited from theme'}",
           f"Font sizes used (pt): {D['font_sizes_pt']}",
           f"Title sizes (pt): {D['title_sizes_pt']}",
           f"Text colors: {D['text_colors']}",
           f"Fill colors: {D['fill_colors']}",
           f"Layouts: {D['layouts']}",
           f"Slides without title placeholder: {D['slides_without_title']}",
           f"Animations on: {D['slides_with_animation']} · transitions on: {D['slides_with_transition']}",
           "",
           "| # | title | body words | min pt | imgs (low-res<150ppi / distorted / no alt) | charts (3D / legend) | tables |",
           "|---|---|---|---|---|---|---|"]
    for r in d["slides"]:
        imgs = r["images"]
        low = sum(1 for i in imgs if i.get("effective_ppi") and i["effective_ppi"] < 150)
        dist = sum(1 for i in imgs if (i.get("aspect_distortion") or 0) > 0.02)
        noalt = sum(1 for i in imgs if not i["alt_text"])
        ch = r["charts"]
        c3d = sum(c["is_3d"] for c in ch)
        cleg = sum(bool(c["has_legend"]) for c in ch)
        tabs = ", ".join(f"{t['rows']}x{t['cols']}" for t in r["tables"])
        title = (r["title"] or "—")[:70]
        out.append(f"| {r['slide']} | {title} | {r['body_words']} | {r['min_size'] or 'inh'} | "
                   f"{len(imgs)} ({low}/{dist}/{noalt}) | {len(ch)} ({c3d}/{cleg}) | {tabs} |")
    return "\n".join(out)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    data = analyze(sys.argv[1])
    print(json.dumps(data, indent=2, default=str) if "--json" in sys.argv else summary(data))
